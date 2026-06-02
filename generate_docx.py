# -*- coding: utf-8 -*-
"""生成李首男简历 .docx 文件"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── 页面设置 ──
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.5
style.font.color.rgb = RGBColor(0x2c, 0x2c, 0x2c)
rPr = style.element.get_or_add_rPr()
rFonts = OxmlElement('w:rFonts')
rFonts.set(qn('w:eastAsia'), '宋体')
rPr.insert(0, rFonts)

def set_run_font(run, name, ea_name, size, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), ea_name)
    rPr.insert(0, rFonts)

def add_heading_styled(text, size=14, space_before=10, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, '黑体', '黑体', size, True, RGBColor(0x8b, 0x6f, 0x4e))
    return p

def add_body(text, size=10.5, bold=False, color=None, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.45
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    c = color or RGBColor(0x4a, 0x4a, 0x4a)
    set_run_font(run, '宋体', '宋体', size, bold, c)
    return p

def add_highlight(text, size=10):
    return add_body(text, size=size, bold=True, color=RGBColor(0x8b, 0x6f, 0x4e))

def add_divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('—' * 45)
    set_run_font(run, '宋体', '宋体', 6, False, RGBColor(0xbb, 0xbb, 0xbb))

# ============================================================
# 头部
# ============================================================
p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.space_before = Pt(0)
run = p.add_run('李 首 男')
set_run_font(run, '宋体', '宋体', 24, True, RGBColor(0x1c, 0x1c, 0x1c))

p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
run = p.add_run('电话: 18241591847 | 邮箱: 13314081391@163.com | 年龄: 25岁 | 共青团员')
set_run_font(run, '宋体', '宋体', 9.5, False, RGBColor(0x66, 0x66, 0x66))

p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run('沈阳师范大学 · 市场营销（本科）| 辽宁沈阳 · 沈北新区')
set_run_font(run, '宋体', '宋体', 9.5, False, RGBColor(0x66, 0x66, 0x66))

add_divider()

# ============================================================
# 在校学习课程
# ============================================================
add_heading_styled('在校学习课程', size=12, space_before=8)

courses = [
    ('管理与经济基础：', '管理学 · 微观经济学 · 宏观经济学 · 统计学'),
    ('营销与销售核心：', '市场营销学 · 销售管理 · 广告学 · 消费者行为学'),
    ('运营与供应链：',   '物流管理 · 会计学基础 · 财务管理'),
]
for label, content in courses:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(label)
    set_run_font(run, '黑体', '黑体', 10, True, RGBColor(0x33, 0x33, 0x33))
    run2 = p.add_run(content)
    set_run_font(run2, '宋体', '宋体', 10, False, RGBColor(0x55, 0x55, 0x55))

# ============================================================
# 校园经历
# ============================================================
add_heading_styled('校园经历', size=12, space_before=10)

add_body('沈阳师范大学 · 市场营销（本科） · 2020.09 — 2026.06', size=10)
add_body('班级团支部书记（2020.10—2024.06）| 团委组织部部员（大一大二）| 团委第二课堂部部长（大三）', size=10, bold=True)
add_body('🏅 三次院级奖学金 | ⭐ 一次校优秀团员', size=10, bold=True, color=RGBColor(0x8b, 0x6f, 0x4e))

campus_items = [
    ('班级团支部书记',
     '统筹并组织了多次大型志愿者活动、爱心献血以及疫情期间的校园志愿服务，累计服务时长超百小时，展现了极强的号召力与公共服务意识。'),
    ('团委组织部部员（大一至大二）',
     '参与校团委组织部的日常事务，负责团员发展、团籍管理及团内活动组织工作，培养了严谨的组织协调能力。'),
    ('团委第二课堂部部长（大三）',
     '核心主导：由本人牵头并协调学校团委老师与技术团队，从 0 到 1 共同开发并上线了"沈师青课堂"APP。深入参与该系统二课考核标准的编纂与顶层设计，将 APP 内的学分考核与学生毕业资质进行深度挂钩，实现了全校学生的高频次活跃与刚性覆盖。'),
    ('多元社会实践',
     '在校期间勇于走出舒适圈，先后尝试了摆摊、后厨、搬家、家教、店员、外卖、进厂等 10 余种跨界兼职工作。积累了极其深厚的底层社会经验，掌握了针对不同阶层与背景人群的跨沟通技巧，培养了极强的抗压能力与同理心。'),
    ('知行合一的深度实践（大四至大六）',
     '在大四于实体企业完成深度实习后，深感校园理论知识需与商业前沿高频碰撞。为此，后续选择主导进行长达两年的"GAP 式"高强度社会实习，主动将市场营销理论在 ToB / ToC 多条商业赛道中进行全面内化。'),
]

for title, desc in campus_items:
    add_body(title, size=10.5, bold=True)
    add_body(desc, size=10, indent=0.4)

# ============================================================
# 实习经历
# ============================================================
add_heading_styled('实习经历', size=12, space_before=10)

companies = [
    ('久屹机械制造有限公司', '市场拓展与技术协调（实习生）', '2023.07 — 2023.10',
     '负责向工业级客户讲解、演示公司机械产品，针对不同客户的生产业务场景进行定制化、差异化的产品推荐。充当市场与生产技术端的纽带，精准翻译并反馈客户的非标定制需求给生产研发部门，跟踪协调技术方案的落地。',
     '了解了工业制造企业的运作架构，切实掌握了企业内部跨部门（销售—生产—技术）交接、协同与复盘流程；提升了将客户痛点转化为产品技术语言的跨界沟通与项目协调能力。',
     None),
    ('罗森（沈阳）', '区域零售门店运营与供应链协同', '2024.01 — 2024.09',
     '深度参与辽宁省内多家直营及加盟门店的日常经营管理，具备多店并发运营的全局视角。运用 ERP 系统实时动态监控商品数据，精细化统计临期损耗与畅销品类，执行高效的跨店转货调拨，优化库存结构。参与规划并落地零售门店的日常促销、会员专属礼品管理以及重大节假日主题活动。',
     '积累了极强的新零售实体店运营经验、面销技巧及敏锐的数据看板分析能力；深刻理解了大型连锁零售业态的商品全生命周期管理、物流转配与供应链跨部门协同机制。',
     None),
    ('博宇金属贸易公司', 'ToB 大客户销售（期货 / 跟单）', '2024.10 — 2025.06',
     '通过行业展会、专业论坛、宏观新闻及产业研报，深度挖掘潜在企业级客户，利用电话销售进行精准触达与初步需求建档。负责 ToB 工业品销售、大宗商品期货交易与全流程跟单；精细化对接客户需求，协调供应链、物流团队。定期进行客户电访与回访，维护长期合作关系。',
     '培养了敏锐的行业宏观经济与期货市场分析能力，能够基于专业数据为客户提供顾问式销售方案；掌握了 ToB 大客户开发流转、商务合同谈判、风险控制及全周期跟单管理的硬核技能。',
     '📊 月平均处理业务额超十万元'),
    ('沈阳天童美语', '营销推广与私域运营（ToC）', '2025.08 — 2025.12',
     '负责沈阳核心商圈及大型幼儿园等高客群密集场所的地推拓客，通过精准人群画像分析，日均高效触达目标家长群体，实现公域流量向私域的转化。负责潜在客户群体的长期维护，通过高价值内容提升社群活跃度与信任度；每周六独立/协助组织线下家长学生联动活动。',
     '沉淀了扎实的一线面销技巧与抗压能力，深刻理解 ToC 消费心理学；掌握了从"线下地推 → 社群沉淀 → 活动裂变 → 深度促活"的全链路私域运营闭环思维。',
     '💰 成功转化学员，实现盈利四万余元'),
    ('锦书教育', '在线教育销售与直播运营', '2026.01 — 2026.02',
     '面向已购低价/体验课学员进行线上直播授课，通过高效的课堂互动与价值锚定，激发学员兴趣，建立讲师信任度。配合直播节点，通过电话对家长进行深度回访与需求诊断，制定针对性续费方案，促成高价正价课的加购与二次转化。',
     '跑通了在线教育"直播带货 / 高转化授课 + 电话精准社群催单"的复合型销售模式；提升了线上公众表达、临场应变能力以及精细化用户生命周期管理（LTV）的意识。',
     '👥 累计覆盖学员 107 人'),
]

for name, role, date, duty, gain, highlight in companies:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(name)
    set_run_font(run, '黑体', '黑体', 10.5, True, RGBColor(0x1c, 0x1c, 0x1c))
    run2 = p.add_run(f'  |  {role}  |  {date}')
    set_run_font(run2, '宋体', '宋体', 9, False, RGBColor(0x88, 0x88, 0x88))

    add_body('工作职责：' + duty, size=10, indent=0.4)
    add_body('收获：' + gain, size=10, indent=0.4)
    if highlight:
        add_highlight(highlight, size=10)

# ============================================================
# 专业技能与证书
# ============================================================
add_heading_styled('专业技能与证书', size=12, space_before=10)

add_body('前沿 AI 技能', size=10.5, bold=True)
add_body('阿里云 ACP 大模型高级工程师认证 — 深入理解大语言模型应用落地、Prompt 工程及大模型赋能业务流程', size=10, indent=0.4)
add_highlight('📊 全国持有此认证者不足五万人', size=10)
add_body('微软认证 Azure AI Fundamentals（AI-900）— 具备国际前沿的云计算及人工智能基础解决方案认知', size=10, indent=0.4)

add_body('综合素养', size=10.5, bold=True)
add_body('快速跨行业学习能力 · 商业数据看板分析能力 · 跨部门复杂项目协同能力', size=10, indent=0.4)

# ============================================================
# 个人特点与爱好
# ============================================================
add_heading_styled('个人特点与爱好', size=12, space_before=10)
hobbies = [
    ('自我驱动与终身学习：',
     '保持高强度的阅读习惯，大学期间在校借阅并研读历史学、社会学、世界系统理论等文科专业书籍百余本。构建了宏观的历史跨度思维、严密的逻辑推演能力以及面对复杂商业问题时的"多视角分析"能力。'),
    ('探索精神与抗压柔韧度：',
     '热爱轻资产探险式"穷游"，曾乘坐绿皮火车足迹遍布从大西北到烟雨江南的 10 个省份。极具韧性，乐于在艰苦环境中寻找最优解，拥有开阔的眼界与胸怀。'),
    ('竞技心态：',
     '乒乓球运动爱好者，具备极高的参与热情与屡败屡战的乐观主义精神。'),
]
for label, text in hobbies:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.45
    run = p.add_run(label)
    set_run_font(run, '黑体', '黑体', 10, True, RGBColor(0x1c, 0x1c, 0x1c))
    run2 = p.add_run(text)
    set_run_font(run2, '宋体', '宋体', 10, False, RGBColor(0x55, 0x55, 0x55))

# ============================================================
# 结尾 — 网站链接
# ============================================================
add_divider()

p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(4)
run = p.add_run('📎 在线简历 & AI 问答助手')
set_run_font(run, '黑体', '黑体', 14, True, RGBColor(0x8b, 0x6f, 0x4e))

p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run('https://snli258.cn')
set_run_font(run, 'Calibri', 'Calibri', 13, True, RGBColor(0x00, 0x66, 0xcc))

p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run('欢迎 HR 或老板访问以上链接，查看完整在线简历，并体验由本人基于')
set_run_font(run, '宋体', '宋体', 10, False, RGBColor(0x99, 0x99, 0x99))
run2 = p.add_run('百炼平台、DeepSeek V4、RAG、Agent 与 Claude Code')
set_run_font(run2, '黑体', '黑体', 10, True, RGBColor(0x8b, 0x6f, 0x4e))
run3 = p.add_run('独立搭建的 AI 问答助手。')
set_run_font(run3, '宋体', '宋体', 10, False, RGBColor(0x99, 0x99, 0x99))

p = doc.add_paragraph()
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('您的任何问题，AI 助手均可为您即时解答。感谢您的关注！')
set_run_font(run, '宋体', '宋体', 10, False, RGBColor(0x99, 0x99, 0x99))

# ── 保存 ──
output_path = r'G:\MyAiProject\简历\李首男简历-新版.docx'
doc.save(output_path)
print('Done: ' + output_path)
print('Size: ' + str(os.path.getsize(output_path)))
